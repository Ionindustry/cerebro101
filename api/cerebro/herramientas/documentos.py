"""Extracción de texto de documentos recibidos (adjuntos de correo): PDF, Word, Excel, CSV, texto y HTML.

Los documentos vienen de terceros y no son de fiar: solo se LEE su texto, nunca se ejecuta nada (ni macros ni scripts), se comprueba que
el contenido sea del tipo que dice el nombre, se limita el tamaño y se desconfía de los ZIP que se descomprimen enormes.
Las imágenes y los PDF escaneados se leen con OCR (ocr.py); si no hay OCR en el servidor se avisa en vez de inventar.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from html.parser import HTMLParser
from xml.etree import ElementTree

from . import ocr as O

OCR_AVISO = "Texto obtenido por OCR: puede tener errores, sobre todo en cifras, fechas y NIF. Comprueba los datos importantes antes de usarlos."

EXTENSIONES = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".html", ".htm", ".xml", ".json", ".md"}
IMAGENES = {".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".bmp", ".webp"}
MAX_DESCOMPRIMIDO = 80 * 1024 * 1024        # límite de lo que puede ocupar un docx/xlsx ya descomprimido (anti «zip bomb»)
PAGINAS_MAX = 80
FILAS_MAX = 2000


class DocumentoNoValido(ValueError):
    """El mensaje es para el modelo y para la persona: dice qué pasa con el documento."""


def extension(nombre: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]{1,5}$", nombre or "")
    return m.group(0).lower() if m else ""


def nombre_seguro(nombre: str) -> str:
    """Solo el nombre del fichero, sin rutas ni caracteres raros (se usa al guardarlo en el ERP)."""
    base = re.split(r"[\\/]", nombre or "")[-1]
    base = re.sub(r"[^\w.\- ]", "_", base, flags=re.UNICODE).strip(" .")
    return (base or "adjunto")[:120]


def _comprobar_zip(datos: bytes) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(datos))
    except zipfile.BadZipFile as e:
        raise DocumentoNoValido("El fichero no es un documento de Office válido (está dañado o no es lo que dice su extensión)") from e
    if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPRIMIDO:
        raise DocumentoNoValido("El documento ocupa demasiado al descomprimirse; no se abre por seguridad")
    return z


def _pdf(datos: bytes) -> tuple[str, dict]:
    if not datos.lstrip()[:5] == b"%PDF-":
        raise DocumentoNoValido("El fichero dice ser un PDF pero no lo es")
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError
    try:
        lector = PdfReader(io.BytesIO(datos))
        if lector.is_encrypted and not lector.decrypt(""):
            raise DocumentoNoValido("El PDF está protegido con contraseña: no se puede leer sin ella")
        paginas = len(lector.pages)
        partes = []
        for i, p in enumerate(lector.pages[:PAGINAS_MAX], 1):
            partes.append(f"[Página {i}]\n{(p.extract_text() or '').strip()}")
    except DocumentoNoValido:
        raise
    except (PyPdfError, ValueError, KeyError, TypeError, OSError, RecursionError) as e:
        raise DocumentoNoValido(f"No se ha podido leer el PDF ({type(e).__name__})") from e
    texto = "\n\n".join(partes)
    info = {"paginas": paginas}
    if paginas > PAGINAS_MAX:
        info["aviso_paginas"] = f"Solo se han leído las {PAGINAS_MAX} primeras páginas de {paginas}"
    if len(re.sub(r"\[Página \d+\]|\s", "", texto)) < 3:
        # sin texto seleccionable: es un escaneo o un PDF de imágenes
        try:
            texto, leidas = O.ocr_pdf(datos)
            info.update(ocr=True, aviso_ocr=OCR_AVISO)
            if paginas > leidas:
                info["aviso_paginas"] = f"Con OCR solo se han leído las {leidas} primeras páginas de {paginas}"
        except O.OcrNoDisponible:
            info["aviso_documento"] = "El PDF no tiene texto seleccionable (parece escaneado) y este servidor no tiene OCR: que lo revise una persona"
        except O.OcrFallido as e:
            info["aviso_documento"] = f"El PDF es un escaneo y el OCR ha fallado: {e}"
    return texto, info


def _docx(datos: bytes) -> tuple[str, dict]:
    z = _comprobar_zip(datos)
    if "word/document.xml" not in z.namelist():
        raise DocumentoNoValido("El fichero no es un documento Word (.docx) válido")
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        raiz = ElementTree.fromstring(z.read("word/document.xml"))
    except ElementTree.ParseError as e:
        raise DocumentoNoValido("El documento Word está dañado") from e
    parrafos = []
    for p in raiz.iter(f"{ns}p"):
        t = "".join(" " if x.tag in (f"{ns}tab", f"{ns}br") else (x.text or "") for x in p.iter() if x.tag in (f"{ns}t", f"{ns}tab", f"{ns}br"))
        if t.strip():
            parrafos.append(t.strip())
    return "\n".join(parrafos), {"aviso_macros": "Se ignoran macros y objetos incrustados"} if "word/vbaProject.bin" in z.namelist() else {}


def _xlsx(datos: bytes) -> tuple[str, dict]:
    _comprobar_zip(datos)
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl lanza varios tipos según el daño
        raise DocumentoNoValido(f"No se ha podido leer el Excel ({type(e).__name__})") from e
    partes, filas = [], 0
    for hoja in libro.worksheets:
        partes.append(f"[Hoja «{hoja.title}»]")
        for fila in hoja.iter_rows(values_only=True):
            if filas >= FILAS_MAX:
                partes.append(f"… (se leen como máximo {FILAS_MAX} filas)")
                break
            if any(c is not None and str(c).strip() for c in fila):
                partes.append(" | ".join("" if c is None else str(c) for c in fila).rstrip(" |"))
                filas += 1
    libro.close()
    return "\n".join(partes), {"hojas": len(libro.worksheets)}


class _Texto(HTMLParser):
    def __init__(self):
        super().__init__()
        self.partes, self._ignorar = [], 0

    def handle_starttag(self, tag, attrs):
        self._ignorar += tag in ("script", "style")
        if tag in ("br", "p", "div", "tr", "li", "h1", "h2", "h3"):
            self.partes.append("\n")

    def handle_endtag(self, tag):
        self._ignorar -= tag in ("script", "style") and self._ignorar > 0

    def handle_data(self, data):
        if not self._ignorar:
            self.partes.append(data)


def _decodificar(datos: bytes) -> str:
    for cod in ("utf-8-sig", "cp1252"):
        try:
            return datos.decode(cod)
        except UnicodeDecodeError:
            continue
    return datos.decode("latin-1", errors="replace")


def _csv(datos: bytes) -> tuple[str, dict]:
    texto = _decodificar(datos)
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = []
    for i, fila in enumerate(csv.reader(io.StringIO(texto), dialecto)):
        if i >= FILAS_MAX:
            filas.append(f"… (se leen como máximo {FILAS_MAX} filas)")
            break
        filas.append(" | ".join(fila))
    return "\n".join(filas), {}


def extraer_texto(nombre: str, datos: bytes) -> tuple[str, dict]:
    """(texto completo, información). Lanza DocumentoNoValido con un mensaje claro si no se puede o no se debe leer."""
    ext = extension(nombre)
    if ext in IMAGENES:
        try:
            return O.ocr_imagen(datos, ext), {"ocr": True, "aviso_ocr": OCR_AVISO}
        except O.OcrNoDisponible:
            return "", {"aviso_documento": "Es una imagen y este servidor no tiene OCR: que la revise una persona"}
        except O.OcrFallido as e:
            raise DocumentoNoValido(str(e)) from e
    if ext not in EXTENSIONES:
        raise DocumentoNoValido(f"Tipo de fichero «{ext or 'desconocido'}» no admitido. Se leen: {', '.join(sorted(EXTENSIONES))}")
    if ext == ".pdf":
        return _pdf(datos)
    if ext == ".docx":
        return _docx(datos)
    if ext == ".xlsx":
        return _xlsx(datos)
    if ext == ".csv":
        return _csv(datos)
    texto = _decodificar(datos)
    if ext in (".html", ".htm"):
        p = _Texto()
        p.feed(texto)
        texto = re.sub(r"\n\s*\n+", "\n", "".join(p.partes))
    return texto.strip(), {}
